"""Export the additional 40-stock, five-factor weekly dataset to Excel.

Input files (created by collect_additional_40_staples.py):
    data/additional_40/additional_40_weekly_5_factors.csv
    data/additional_40/additional_40_universe.csv
    data/additional_40/additional_40_daily_ohlcv.csv

Output file:
    data/additional_40/additional_40_five_factors_targets_final.xlsx

The source CSV already contains the forward one-week target calculated from
adjusted closing prices:

    target1 = ln(AdjustedClose_week_end(t + 1) /
                 AdjustedClose_week_end(t))

This script renames ``target`` to ``target1`` and recreates ``target2`` using
the agreed classification thresholds.  It deliberately keeps the source
target rather than recalculating it after filtering rows: the source may hold
the final label using a subsequent price observation that is not itself part
of the modelling sample.

``buy_price`` is the first unadjusted open price of the following week;
``sell_price`` is the last unadjusted close price of that same week.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


REQUIRED_SOURCE_COLUMNS = [
    "date",
    "ticker",
    "revenue_growth_yoy",
    "operating_profitability",
    "cashflow_to_assets",
    "debt_to_assets",
    "asset_growth_yoy",
    "target",
]

OUTPUT_COLUMNS = [
    "date",
    "ticker",
    "stock_name",
    "buy_price",
    "sell_price",
    "revenue_growth_yoy",
    "operating_profitability",
    "cashflow_to_assets",
    "debt_to_assets",
    "asset_growth_yoy",
    "target1",
    "target2",
]

REQUIRED_DAILY_COLUMNS = ["date", "ticker", "open", "close"]


def classify_target(target: pd.Series) -> pd.Series:
    """Convert target1 to the five-class target2 label.

    0: target1 < -3%; 1: [-3%, -1%); 2: [-1%, +1%];
    3: (+1%, +3%]; 4: target1 > +3%.
    """
    labels = np.select(
        [
            target < -0.03,
            (target >= -0.03) & (target < -0.01),
            (target >= -0.01) & (target <= 0.01),
            (target > 0.01) & (target <= 0.03),
            target > 0.03,
        ],
        [0, 1, 2, 3, 4],
        default=-1,
    )
    if (labels == -1).any():
        raise ValueError("target1 contains missing or unclassifiable values.")
    return pd.Series(labels, index=target.index, dtype="int8")


def add_next_week_prices(weekly: pd.DataFrame, daily_path: Path) -> pd.DataFrame:
    """Attach next week's first open and last close to each feature week."""
    daily = pd.read_csv(daily_path, parse_dates=["date"])
    missing_columns = set(REQUIRED_DAILY_COLUMNS) - set(daily.columns)
    if missing_columns:
        raise ValueError(f"Daily input is missing columns: {sorted(missing_columns)}")
    if daily.duplicated(["date", "ticker"]).any():
        raise ValueError("Daily input has duplicate date-ticker keys.")
    if daily[REQUIRED_DAILY_COLUMNS].isna().any().any():
        raise ValueError("Daily input has missing values in date, ticker, open, or close.")

    daily = daily.sort_values(["ticker", "date"]).copy()
    daily["week"] = daily["date"].dt.to_period("W-FRI")
    weekly_prices = (
        daily.groupby(["ticker", "week"], as_index=False)
        .agg(
            week_last_date=("date", "last"),
            week_first_open=("open", "first"),
            week_last_close=("close", "last"),
        )
        .sort_values(["ticker", "week_last_date"])
    )
    weekly_prices["buy_price"] = weekly_prices.groupby("ticker")["week_first_open"].shift(-1)
    weekly_prices["sell_price"] = weekly_prices.groupby("ticker")["week_last_close"].shift(-1)

    merged = weekly.merge(
        weekly_prices[["ticker", "week_last_date", "buy_price", "sell_price"]],
        left_on=["ticker", "date"],
        right_on=["ticker", "week_last_date"],
        how="left",
        validate="one_to_one",
    ).drop(columns="week_last_date")
    if merged[["buy_price", "sell_price"]].isna().any().any():
        raise ValueError("Some weekly rows cannot be matched to complete next-week prices.")
    return merged


def build_export(weekly_path: Path, universe_path: Path, daily_path: Path) -> pd.DataFrame:
    """Load, validate, and select the fields for the final Excel file."""
    weekly = pd.read_csv(weekly_path, parse_dates=["date"])
    universe = pd.read_csv(universe_path, usecols=["ticker", "company_name"])

    missing_columns = set(REQUIRED_SOURCE_COLUMNS) - set(weekly.columns)
    if missing_columns:
        raise ValueError(f"Weekly input is missing columns: {sorted(missing_columns)}")
    if universe["ticker"].duplicated().any():
        raise ValueError("The universe file has duplicate ticker symbols.")
    if weekly.duplicated(["date", "ticker"]).any():
        raise ValueError("Weekly input has duplicate date-ticker keys.")
    if weekly[REQUIRED_SOURCE_COLUMNS].isna().any().any():
        counts = weekly[REQUIRED_SOURCE_COLUMNS].isna().sum()
        raise ValueError(f"Weekly input has missing values:\n{counts[counts > 0]}")

    result = add_next_week_prices(weekly[REQUIRED_SOURCE_COLUMNS].copy(), daily_path)
    result = result.merge(
        universe.rename(columns={"company_name": "stock_name"}),
        on="ticker",
        how="left",
        validate="many_to_one",
    )
    if result["stock_name"].isna().any():
        unknown = sorted(result.loc[result["stock_name"].isna(), "ticker"].unique())
        raise ValueError(f"Ticker(s) missing from universe file: {unknown}")

    result = result.rename(columns={"target": "target1"})
    result["target2"] = classify_target(result["target1"])
    result = result[OUTPUT_COLUMNS].sort_values(["date", "ticker"]).reset_index(drop=True)

    if result.isna().any().any():
        raise ValueError("The final export has missing values.")
    return result


def write_excel(data: pd.DataFrame, output_path: Path) -> None:
    """Write an analysis-ready Excel workbook with basic, readable formatting."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl", date_format="yyyy-mm-dd") as writer:
        data.to_excel(writer, index=False, sheet_name="weekly_five_factors")
        sheet = writer.book["weekly_five_factors"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        header_fill = PatternFill("solid", fgColor="C0704B")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill

        for row in sheet.iter_rows(min_row=2, min_col=1, max_col=1):
            row[0].number_format = "yyyy-mm-dd"
        for col in range(4, 6):
            for row in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
                row[0].number_format = "0.00"
        for col in range(6, 12):
            for row in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
                row[0].number_format = "0.000000"
        for row in sheet.iter_rows(min_row=2, min_col=12, max_col=12):
            row[0].number_format = "0"

        widths = {
            1: 13, 2: 10, 3: 38, 4: 14, 5: 14, 6: 20, 7: 23,
            8: 20, 9: 17, 10: 18, 11: 14, 12: 10,
        }
        for column, width in widths.items():
            sheet.column_dimensions[get_column_letter(column)].width = width


def main() -> None:
    project_dir = Path(__file__).resolve().parent
    default_data_dir = project_dir / "data" / "additional_40"

    parser = argparse.ArgumentParser(description="Export five factors and two targets for 40 additional stocks.")
    parser.add_argument("--input", type=Path, default=default_data_dir / "additional_40_weekly_5_factors.csv")
    parser.add_argument("--universe", type=Path, default=default_data_dir / "additional_40_universe.csv")
    parser.add_argument("--daily", type=Path, default=default_data_dir / "additional_40_daily_ohlcv.csv")
    parser.add_argument("--output", type=Path, default=default_data_dir / "additional_40_five_factors_targets_final.xlsx")
    args = parser.parse_args()

    if not args.input.exists() or not args.universe.exists() or not args.daily.exists():
        raise FileNotFoundError("Input CSV file(s) are missing. Run collect_additional_40_staples.py first.")

    result = build_export(args.input, args.universe, args.daily)
    try:
        write_excel(result, args.output)
    except PermissionError as error:
        raise PermissionError(
            f"Cannot overwrite {args.output}. Close the workbook in Excel and run the script again."
        ) from error
    print(f"Created: {args.output}")
    print(f"Rows: {len(result):,}; stocks: {result['ticker'].nunique()}; columns: {len(result.columns)}")
    print(f"Date range: {result['date'].min():%Y-%m-%d} to {result['date'].max():%Y-%m-%d}")


if __name__ == "__main__":
    main()
